#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
edit_report_ppt.py
------------------------------------------------------------------
기존 결재본 PPTX에 하우스 스타일 슬라이드를 '추가'하는 도구.

새로 생성하지 않고 원본 파일을 열어 슬라이드를 삽입하므로,
사용자가 직접 수정한 위치·문구·서식이 그대로 보존된다.

  · 표 서식(머리행 002060 / 테두리 BFBFBF 0.75pt / 나눔바른고딕)
  · [[파랑]] {{파랑+노랑형광}} 마크업
  · 하단 전폭 의견 바(FFFF99)
을 build_report_ppt.js 와 동일하게 재현한다.
------------------------------------------------------------------
"""
import copy
import re

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS = {"a": A}

FONT = "나눔바른고딕"
TBL_HEAD = "002060"
BAR = "FFFF99"
BLUE = "0000FF"
MARK = "FFFF00"
TEXT = "1F1F1F"
LINE = "BFBFBF"

MX, MW, MTOP, HEAD_H = 0.140, 10.472, 0.157, 0.337


# ── 마크업 ─────────────────────────────────────────────────────
def parse(src):
    """[[파랑]] / {{파랑+형광}} / \n 을 (text, color, highlight, br) 목록으로."""
    out = []
    for li, line in enumerate(str(src).split("\n")):
        parts = [s for s in re.split(r"(\[\[[^\]]*\]\]|\{\{[^}]*\}\})", line) if s != ""]
        if not parts:
            parts = [""]
        for si, seg in enumerate(parts):
            color, hl = None, None
            t = seg
            if seg.startswith("[["):
                t, color = seg[2:-2], BLUE
            elif seg.startswith("{{"):
                t, color, hl = seg[2:-2], BLUE, MARK
            out.append((t, color, hl, si == len(parts) - 1))
        out[-1] = (out[-1][0], out[-1][1], out[-1][2], True)
    return out


def plain(s):
    return re.sub(r"\[\[|\]\]|\{\{|\}\}", "", str(s))


def em_width(s):
    w = 0.0
    for ch in str(s):
        w += 1.0 if (0x1100 <= ord(ch) <= 0x11FF or 0x3000 <= ord(ch) <= 0x303F
                     or 0x3130 <= ord(ch) <= 0x318F or 0xAC00 <= ord(ch) <= 0xD7AF
                     or 0xFF00 <= ord(ch) <= 0xFFEF) else 0.52
    return w


def est_lines(text, width_in, pt, pad=0.19):
    usable = max(0.4, width_in - pad)
    per = max(4.0, usable * 72 / pt)
    n = 0
    for seg in plain(text).split("\n"):
        n += max(1, int(-(-em_width(seg) // per)))
    return n


def _set_runs(tf, src, pt, bold=True, color=TEXT, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.MIDDLE, spacing=None):
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    paras = []
    cur = tf.paragraphs[0]
    paras.append(cur)
    for text, col, hl, br in parse(src):
        r = cur.add_run()
        r.text = text
        f = r.font
        f.name, f.size, f.bold = FONT, Pt(pt), bold
        f.color.rgb = RGBColor.from_string(col or color)
        rPr = r._r.get_or_add_rPr()
        for tag in ("ea", "cs"):
            el = rPr.makeelement("{%s}%s" % (A, tag), {"typeface": FONT})
            rPr.append(el)
        if hl:
            h = rPr.makeelement("{%s}highlight" % A, {})
            c = rPr.makeelement("{%s}srgbClr" % A, {"val": hl})
            h.append(c)
            rPr.insert(0, h)
        if br:
            new = copy.deepcopy(cur._p)
            for child in list(new):
                if child.tag.endswith("}r") or child.tag.endswith("}br"):
                    new.remove(child)
            cur._p.addnext(new)
            from pptx.text.text import _Paragraph
            cur = _Paragraph(new, tf)
            paras.append(cur)
    # 마지막 빈 문단 제거
    if len(paras) > 1 and not paras[-1].runs:
        paras[-1]._p.getparent().remove(paras[-1]._p)
        paras.pop()
    for p in paras:
        p.alignment = align
        if spacing:
            p.line_spacing = spacing


# ── 요소 ───────────────────────────────────────────────────────
def add_heading(slide, text, y, w=7.2):
    tb = slide.shapes.add_textbox(Inches(MX), Inches(y), Inches(w), Inches(HEAD_H))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    _set_runs(tf, text, 14, True, TEXT, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    return y + HEAD_H + 0.08


def _border(tcPr, color=LINE, w=9525):
    for tag in ("lnL", "lnR", "lnT", "lnB"):
        old = tcPr.find("{%s}%s" % (A, tag))
        if old is not None:
            tcPr.remove(old)
        ln = tcPr.makeelement("{%s}%s" % (A, tag),
                              {"w": str(w), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
        fill = tcPr.makeelement("{%s}solidFill" % A, {})
        clr = tcPr.makeelement("{%s}srgbClr" % A, {"val": color})
        fill.append(clr)
        ln.append(fill)
        tcPr.insert(0, ln)


def add_table(slide, y, columns, rows, col_w, pt=10, bottom=7.28, grow=True):
    n = len(rows)
    avail = bottom - y - 0.10
    while True:
        lh = pt * 1.34 / 72
        row_h = [max(0.30, max(est_lines(c, col_w[i], pt) for i, c in enumerate(r)) * lh + 0.13)
                 for r in rows]
        need = sum(row_h) + 0.26
        if need <= avail or pt <= 8:
            break
        pt -= 0.5
    if grow and need < avail:
        extra = min((avail - need) / n, 0.52)
        row_h = [h + extra for h in row_h]
        need += extra * n

    gf = slide.shapes.add_table(n + 1, len(columns), Inches(MX), Inches(y),
                                Inches(MW), Inches(need))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    for i, w in enumerate(col_w):
        tbl.columns[i].width = Inches(w)
    tbl.rows[0].height = Inches(0.26)
    for i, h in enumerate(row_h):
        tbl.rows[i + 1].height = Inches(h)

    for ci, name in enumerate(columns):
        c = tbl.cell(0, ci)
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor.from_string(TBL_HEAD)
        c.margin_left = c.margin_right = Pt(7)
        c.margin_top = c.margin_bottom = Pt(2)
        _set_runs(c.text_frame, name, pt, True, "FFFFFF", PP_ALIGN.CENTER)
        _border(c._tc.find("{%s}tcPr" % A))

    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tbl.cell(ri + 1, ci)
            c.fill.background()
            c.margin_left = c.margin_right = Pt(7)
            c.margin_top = c.margin_bottom = Pt(2)
            _set_runs(c.text_frame, val, pt, True, TEXT,
                      PP_ALIGN.CENTER if ci == 0 else PP_ALIGN.LEFT)
            _border(c._tc.find("{%s}tcPr" % A))
    return y + need + 0.22


def add_bar(slide, lines, pt=11):
    n = sum(est_lines("- " + t, 10.8333 - 0.9, pt, 0.1) for t in lines)
    h = max(0.68, n * pt * 1.35 / 72 + 0.22)
    y = 7.5 - h - 0.06
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(y),
                                Inches(10.8333), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(BAR)
    sh.line.color.rgb = RGBColor.from_string(BAR)
    sh.line.width = Pt(0.5)
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    tb = slide.shapes.add_textbox(Inches(0.25), Inches(y), Inches(10.3333), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    _set_runs(tf, "\n".join("- " + t for t in lines), pt, True, TEXT,
              PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)
    return y


def add_note(slide, text, y, pt=8.5):
    n = est_lines(text, MW, pt, 0.1)
    h = n * pt * 1.45 / 72 + 0.04
    tb = slide.shapes.add_textbox(Inches(MX + 0.02), Inches(y), Inches(MW), Inches(h))
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    _set_runs(tf, text, pt, False, "595959", PP_ALIGN.LEFT, MSO_ANCHOR.TOP)
    return y + h + 0.10


def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_masters[0].slide_layouts[0])


def move_slide(prs, old_index, new_index):
    lst = prs.slides._sldIdLst
    ids = list(lst)
    lst.remove(ids[old_index])
    lst.insert(new_index, ids[old_index])
