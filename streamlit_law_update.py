import streamlit as st
import pandas as pd
import sqlite3
import io
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

DEFAULT_DEPT = "영업계획팀"

# ====================== DB ======================
conn = sqlite3.connect("law_updates.db", check_same_thread=False)
cursor = conn.cursor()

cursor.executescript('''
    CREATE TABLE IF NOT EXISTS law_updates (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        law_name TEXT,
        related_dept TEXT,
        source_org TEXT,
        revision_reason TEXT,
        before_content TEXT,
        after_content TEXT,
        effective_date TEXT,
        press_title TEXT,
        press_content TEXT,
        press_date TEXT,
        note TEXT,
        registered_at TEXT
    );
''')
conn.commit()

st.set_page_config(page_title="법규 최신화 프로그램", layout="wide", page_icon="📜")
st.title("📜 법규 최신화 프로그램")
st.caption("법규 개정사항(이유·개정전·개정후) 및 보도자료를 등록하면, 월별 행정절차·법규 변경 보고서(엑셀)와 보도자료 정리 PPT를 자동 생성합니다.")

menu = st.sidebar.selectbox(
    "📋 메뉴 선택",
    ["1. 법규 개정사항 등록",
     "2. 전체 데이터 관리",
     "3. 📊 보고서 생성"]
)

# ====================== 1. 법규 개정사항 등록 ======================
if menu == "1. 법규 개정사항 등록":
    st.subheader("1. 법규 개정사항 등록")

    with st.form("law_update_form", clear_on_submit=True):
        col1, col2, col3 = st.columns(3)
        with col1:
            law_name = st.text_input("법규명 *", placeholder="예: 주택법 시행령")
        with col2:
            related_dept = st.text_input("관련 부서", DEFAULT_DEPT)
        with col3:
            source_org = st.text_input("소관부처/공포기관", placeholder="예: 국토교통부")

        st.markdown("**① 법규 개정 내용**")
        revision_reason = st.text_area("개정 이유", height=90, placeholder="개정 배경 및 사유를 입력하세요.")
        col4, col5 = st.columns(2)
        with col4:
            before_content = st.text_area("개정 전", height=140, placeholder="개정 전 조문/내용을 입력하세요.")
        with col5:
            after_content = st.text_area("개정 후", height=140, placeholder="개정 후 조문/내용을 입력하세요.")
        effective_date = st.date_input("시행일", value=date.today())

        st.markdown("**② 보도자료**")
        press_title = st.text_input("보도자료 제목", placeholder="예: OOO법 개정, 언제부터 시행되나")
        press_content = st.text_area("보도자료 내용", height=140, placeholder="보도자료 핵심 내용을 입력하세요.")
        press_date = st.date_input("보도자료 배포일", value=date.today())

        note = st.text_input("비고")

        submitted = st.form_submit_button("✅ 등록", type="primary", use_container_width=True)
        if submitted:
            if not law_name:
                st.error("법규명은 필수 입력입니다.")
            else:
                cursor.execute('''INSERT INTO law_updates
                    (law_name, related_dept, source_org, revision_reason, before_content, after_content,
                     effective_date, press_title, press_content, press_date, note, registered_at)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?)''',
                    (law_name, related_dept or DEFAULT_DEPT, source_org, revision_reason,
                     before_content, after_content, effective_date.strftime("%Y-%m-%d"),
                     press_title, press_content, press_date.strftime("%Y-%m-%d"), note,
                     datetime.now().strftime("%Y-%m-%d")))
                conn.commit()
                st.success(f"'{law_name}' 개정사항이 등록되었습니다.")

# ====================== 2. 전체 데이터 관리 ======================
elif menu == "2. 전체 데이터 관리":
    st.subheader("2. 전체 데이터 관리")
    df = pd.read_sql('''SELECT id, law_name as 법규명, related_dept as 관련부서, source_org as 소관부처,
                         effective_date as 시행일, registered_at as 등록일, revision_reason as 개정이유,
                         before_content as 개정전, after_content as 개정후,
                         press_title as 보도자료제목, press_date as 보도자료일, note as 비고
                         FROM law_updates ORDER BY id DESC''', conn)
    st.dataframe(df, use_container_width=True)

    if not df.empty:
        del_id = st.number_input("삭제할 ID", 1, step=1, key="del_law")
        if st.button("🗑 삭제"):
            cursor.execute("DELETE FROM law_updates WHERE id=?", (del_id,))
            conn.commit()
            st.success("삭제 완료")
            st.rerun()
    else:
        st.info("등록된 법규 개정사항이 없습니다.")

# ====================== 3. 보고서 생성 ======================
elif menu == "3. 📊 보고서 생성":
    st.subheader("3. 📊 보고서 생성")

    target_month = st.text_input("대상 월 (YYYY-MM)", datetime.now().strftime("%Y-%m"))
    month_basis = st.radio("월 기준", ["등록일 기준", "시행일 기준"], horizontal=True)
    basis_col = "registered_at" if month_basis == "등록일 기준" else "effective_date"

    df = pd.read_sql(f'''SELECT * FROM law_updates
                          WHERE substr({basis_col}, 1, 7) = ?
                          ORDER BY effective_date''', conn, params=(target_month,))

    st.write(f"**{target_month} 대상 건수: {len(df)}건**")
    st.dataframe(
        df[["law_name", "related_dept", "source_org", "revision_reason", "effective_date"]]
        .rename(columns={"law_name": "법규명", "related_dept": "관련부서", "source_org": "소관부처",
                          "revision_reason": "개정이유", "effective_date": "시행일"}),
        use_container_width=True
    )

    col1, col2 = st.columns(2)

    # -------- 엑셀: 월 행정절차·법규 변경 보고사항 --------
    with col1:
        if st.button("📗 엑셀 보고서 생성", type="primary", use_container_width=True):
            if df.empty:
                st.warning("대상 기간에 등록된 법규 개정사항이 없습니다.")
            else:
                wb = Workbook()
                ws = wb.active
                ws.title = "법규변경보고"

                title_text = f"{target_month} 행정절차·법규 변경 보고사항 ({DEFAULT_DEPT})"
                ws.merge_cells("A1:I1")
                ws["A1"] = title_text
                ws["A1"].font = Font(size=14, bold=True)
                ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
                ws.row_dimensions[1].height = 28

                headers = ["순번", "법규명", "관련부서", "소관부처", "개정이유", "개정전", "개정후", "시행일", "비고"]
                ws.append([])
                ws.append(headers)
                header_row = 3
                header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
                header_font = Font(color="FFFFFF", bold=True)
                thin = Side(style="thin", color="B7B7B7")
                border = Border(left=thin, right=thin, top=thin, bottom=thin)
                for c in range(1, len(headers) + 1):
                    cell = ws.cell(row=header_row, column=c)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    cell.border = border

                for i, row in enumerate(df.itertuples(), start=1):
                    r = header_row + i
                    values = [i, row.law_name, row.related_dept, row.source_org, row.revision_reason,
                              row.before_content, row.after_content, row.effective_date, row.note]
                    for c, v in enumerate(values, start=1):
                        cell = ws.cell(row=r, column=c, value=v)
                        cell.alignment = Alignment(horizontal="center" if c in (1, 3, 4, 8) else "left",
                                                    vertical="center", wrap_text=True)
                        cell.border = border

                widths = [6, 22, 12, 14, 26, 30, 30, 12, 16]
                for c, w in enumerate(widths, start=1):
                    ws.column_dimensions[get_column_letter(c)].width = w

                buf = io.BytesIO()
                wb.save(buf)
                buf.seek(0)
                st.download_button(
                    "⬇️ 엑셀 파일 다운로드",
                    data=buf,
                    file_name=f"{target_month}_행정절차_법규변경_보고사항_{DEFAULT_DEPT}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
                st.success("엑셀 보고서가 생성되었습니다.")

    # -------- PPT: 보도자료 정리본 --------
    with col2:
        if st.button("📙 보도자료 정리 PPT 생성", type="primary", use_container_width=True):
            if df.empty:
                st.warning("대상 기간에 등록된 법규 개정사항이 없습니다.")
            else:
                prs = Presentation()
                prs.slide_width = Inches(13.333)
                prs.slide_height = Inches(7.5)

                # 표지
                title_slide = prs.slides.add_slide(prs.slide_layouts[0])
                title_slide.shapes.title.text = f"{target_month} 법규 개정 보도자료 정리"
                title_slide.placeholders[1].text = f"{DEFAULT_DEPT}  |  작성일: {datetime.now().strftime('%Y-%m-%d')}"

                blank_layout = prs.slide_layouts[6]

                for row in df.itertuples():
                    slide = prs.slides.add_slide(blank_layout)

                    # 헤더 바
                    bar = slide.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(1.0))
                    bar.fill.solid()
                    bar.fill.fore_color.rgb = RGBColor(0x30, 0x54, 0x96)
                    bar.line.fill.background()
                    tf = bar.text_frame
                    tf.margin_left = Inches(0.3)
                    tf.word_wrap = True
                    p = tf.paragraphs[0]
                    p.text = row.press_title or row.law_name
                    p.font.size = Pt(24)
                    p.font.bold = True
                    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

                    body = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(12.3), Inches(5.7))
                    btf = body.text_frame
                    btf.word_wrap = True

                    def add_line(label, value, first=False, size=16, bold_label=True):
                        para = btf.paragraphs[0] if first else btf.add_paragraph()
                        run_label = para.add_run()
                        run_label.text = f"{label}: "
                        run_label.font.bold = True
                        run_label.font.size = Pt(size)
                        run_value = para.add_run()
                        run_value.text = value or "-"
                        run_value.font.size = Pt(size)
                        para.space_after = Pt(10)

                    add_line("법규명", row.law_name, first=True)
                    add_line("소관부처", row.source_org)
                    add_line("시행일", row.effective_date)
                    add_line("개정 이유", row.revision_reason)
                    add_line("개정 전 → 개정 후", f"{row.before_content}  →  {row.after_content}")
                    add_line("보도자료 내용", row.press_content)
                    add_line("보도자료 배포일", row.press_date)

                buf = io.BytesIO()
                prs.save(buf)
                buf.seek(0)
                st.download_button(
                    "⬇️ PPT 파일 다운로드",
                    data=buf,
                    file_name=f"{target_month}_법규개정_보도자료_정리본.pptx",
                    mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    use_container_width=True
                )
                st.success("보도자료 정리 PPT가 생성되었습니다.")

st.sidebar.info("① 법규 개정사항 등록 → ② 전체 데이터 관리 → ③ 보고서 생성 순으로 사용하세요.")
